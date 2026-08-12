from operator import truediv
# https://www.youtube.com/watch?v=_uQrJ0TkZlc
print('Hello World')
print("o----")
print(" ||||")
print("*" * 10)

price = 10
print(price)
full_name = "John Smith"
age = 20
is_new = True
print(full_name)
# name = input('what is you name?')
# print("Hi " + name)
# birthyear = input(" birthyear: ")
# age = 2026-int(birthyear)
# print(age)

first = "John"
last = "Smith"
message = first + " [" + last + "] " + "is a coder"
msg = f'{first} [{last}] is a coder'
print(msg)
course = "Python for beginners"
print(len(course))

print(course.upper())
print(course.find('on'))
print(course.replace('Python', 'Java'))
print('Python' in course)


print(10/3)
print(10//3)
print(10%3)
print(10**3)

x = 10
# x = x + 3
x += 3
print(x)

import math

print(math.ceil(5.14))
print(math.floor(5.14))

is_hot = False

if is_hot:
    print("Its a hot day")
else:
    print("Its a cold day")
print("enjoy your day")

hasGoodCredit = True
hasCriminalRecord = True

if hasGoodCredit and not hasCriminalRecord:
    print("Eligible for loan")
else:
    print("Not eligible for loan")

temperature = 30
if temperature > 30:
    print("Temperature is greater than 30")
else:
    print("Temperature is less or equal to 30")

name = "Jayy"
if len(name) < 3:
    print("Name is too short")
elif len(name) > 50:
    print("Name is too long")
else:
    print("Name is ok")


# for i in range(5):
#     print(i)
#
# numbers = [1, 2, 3, 4, 5]
# for number in numbers:
#     # print(number * 'X')
#     output = ""
#     for count in range(number):
#         output += 'x'
#     print(output)

# try:
#     age = int(input("Enter your age: "))
#     income = 2000
#     value = income/ag
#     print(age)
# except ValueError:
#     print("Invalid input")
# except ZeroDivisionError:
#     print("Age cannot be zero")

# Defining classes
class Person:
    def __init__(self, name):
        self.name = name

    def talk(self):
        print(f"Hello {self.name}!")


man = Person("Manfasdf")
print(man.name)
man.talk()


import openpyxl as xl
from openpyxl.chart import BarChart, Reference

def process_workbook(filename):
    wb = xl.load_workbook(filename)
    sheet = wb['Sheet1']
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=3)
        corrected_value = cell.value * 0.9
        corrected_price_cell = sheet.cell(row=row, column=4)
        corrected_price_cell.value = corrected_value

    values = Reference(sheet,
                       min_row=2,
                       max_row=sheet.max_row,
                       min_col=4,
                       max_col=4)

    chart = BarChart()
    chart.add_data(values)
    sheet.add_chart(chart, 'e2')

    wb.save(filename)





